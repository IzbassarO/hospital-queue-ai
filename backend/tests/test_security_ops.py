"""Access control (roles, key lifecycle), audit trail (access log, decision attribution), idempotent decisions,
exports, and the API-gap fields (/config, /me, alert filters)."""

import io
import uuid

import httpx
import pytest
from openpyxl import load_workbook

from conftest import API, KEY_LABEL_PREFIX, TEST_ACTOR_PREFIX, auth

pytestmark = pytest.mark.anyio


# ------------------------------------------------------------------------------------------ authentication
@pytest.mark.parametrize("path", ["/health", f"{API}/health"])
async def test_health_is_open(anon_client, path):
    assert (await anon_client.get(path)).status_code == 200


@pytest.mark.parametrize(
    "method,path",
    [
        ("GET", "/overview"),
        ("GET", "/config"),
        ("GET", "/me"),
        ("GET", "/models"),
        ("GET", "/alerts"),
        ("GET", "/decisions"),
        ("POST", "/decisions"),
        ("GET", "/admin/keys"),
        ("GET", "/admin/access-log"),
    ],
)
async def test_every_other_endpoint_requires_a_key(anon_client, method, path):
    response = await anon_client.request(method, f"{API}{path}", json={} if method == "POST" else None)
    assert response.status_code == 401, response.text
    assert response.json()["detail"] == "missing API key"
    assert response.headers.get("www-authenticate") == "X-API-Key"
    wrong = await anon_client.request(method, f"{API}{path}", headers=auth("hqai_not-a-real-key-000000000000000000"))
    assert wrong.status_code == 401 and wrong.json()["detail"] == "invalid or revoked API key"


def decision_payload(high_load: dict, **extra) -> dict:
    return {
        "region_code": high_load["region_code"],
        "org_code": high_load["org_code"],
        "profile_code": high_load["profile_code"],
        "action": "defer",
        "actor": f"{TEST_ACTOR_PREFIX}{uuid.uuid4().hex[:8]}",
        **extra,
    }


async def test_role_matrix(anon_client, api_keys, high_load, created_decision_ids):
    """viewer: every GET; specialist: + POST /decisions; admin: + key management and the access log."""
    expectations = {
        # (method, path): minimum role
        ("GET", "/overview"): "viewer",
        ("GET", f"/hospitals/{high_load['org_code']}/profiles/{high_load['profile_code']}"): "viewer",
        ("GET", "/decisions"): "viewer",
        ("POST", "/decisions"): "specialist",
        ("GET", "/admin/keys"): "admin",
        ("GET", "/admin/access-log"): "admin",
    }
    rank = {"viewer": 0, "specialist": 1, "admin": 2}
    for role, key in api_keys.items():
        me = (await anon_client.get(f"{API}/me", headers=auth(key))).json()
        assert me["role"] == role and me["label"] == f"{KEY_LABEL_PREFIX}{role}" and me["role_label"]
        for (method, path), minimum in expectations.items():
            body = decision_payload(high_load) if method == "POST" else None
            response = await anon_client.request(method, f"{API}{path}", headers=auth(key), json=body)
            allowed = rank[role] >= rank[minimum]
            expected = (201 if method == "POST" else 200) if allowed else 403
            assert response.status_code == expected, (role, method, path, response.text)
            if method == "POST" and allowed:
                created_decision_ids.append(response.json()["id"])
                assert response.json()["api_key_label"] == f"{KEY_LABEL_PREFIX}{role}"


async def test_admin_creates_lists_and_revokes_keys(anon_client, api_keys):
    admin = auth(api_keys["admin"])
    label = f"{KEY_LABEL_PREFIX}created-viewer"
    created = await anon_client.post(f"{API}/admin/keys", headers=admin, json={"role": "viewer", "label": label})
    assert created.status_code == 201, created.text
    body = created.json()
    key = body["key"]
    assert key.startswith("hqai_") and len(key) >= 40 and key.startswith(body["key_prefix"])

    listed = (await anon_client.get(f"{API}/admin/keys", headers=admin)).json()
    row = next(k for k in listed if k["id"] == body["id"])
    assert row["role"] == "viewer" and row["revoked_at"] is None
    assert all("key" not in k and "key_hash" not in k for k in listed), "keys and hashes are never listed"

    assert (await anon_client.get(f"{API}/overview", headers=auth(key))).status_code == 200
    revoked = await anon_client.post(f"{API}/admin/keys/{body['id']}/revoke", headers=admin)
    assert revoked.status_code == 200 and revoked.json()["revoked_at"]
    assert (await anon_client.get(f"{API}/overview", headers=auth(key))).status_code == 401
    assert (await anon_client.post(f"{API}/admin/keys/999999999/revoke", headers=admin)).status_code == 404
    bad_role = await anon_client.post(f"{API}/admin/keys", headers=admin, json={"role": "root", "label": label})
    assert bad_role.status_code == 422


# ------------------------------------------------------------------------------------------ audit trail
async def test_access_log_records_requests(anon_client, api_keys, high_load):
    marker = f"/regions/{high_load['region_code']}"
    viewer_label = f"{KEY_LABEL_PREFIX}viewer"
    assert (await anon_client.get(f"{API}{marker}", headers=auth(api_keys["viewer"]))).status_code == 200
    assert (await anon_client.get(f"{API}/overview")).status_code == 401

    log = await anon_client.get(
        f"{API}/admin/access-log", headers=auth(api_keys["admin"]), params={"key_label": viewer_label, "limit": 50}
    )
    assert log.status_code == 200
    page = log.json()
    assert set(page) == {"items", "total", "limit", "offset"}
    row = next(r for r in page["items"] if r["path"] == f"{API}{marker}")
    assert row["role"] == "viewer" and row["method"] == "GET" and row["status"] == 200
    assert row["latency_ms"] > 0 and row["ts"] and row["client_ip"]

    unauthenticated = await anon_client.get(
        f"{API}/admin/access-log", headers=auth(api_keys["admin"]), params={"status": 401, "limit": 20}
    )
    assert any(r["key_label"] is None and r["role"] is None for r in unauthenticated.json()["items"])


# ------------------------------------------------------------------------------------------ decisions
async def test_decision_idempotency_key(client, high_load, created_decision_ids):
    key = f"pytest-{uuid.uuid4()}"
    payload = decision_payload(high_load, idempotency_key=key, comment="первая отправка")
    first = await client.post(f"{API}/decisions", json=payload)
    assert first.status_code == 201, first.text
    created_decision_ids.append(first.json()["id"])

    retry = await client.post(f"{API}/decisions", json=payload)
    assert retry.status_code == 200, retry.text
    assert retry.json()["id"] == first.json()["id"] and retry.json()["idempotency_key"] == key

    listed = await client.get(
        f"{API}/decisions", params={"org": high_load["org_code"], "profile": high_load["profile_code"], "limit": 500}
    )
    assert sum(1 for d in listed.json()["items"] if d["idempotency_key"] == key) == 1

    conflict = await client.post(f"{API}/decisions", json={**payload, "action": "reject"})
    assert conflict.status_code == 409 and "idempotency_key" in conflict.json()["detail"]
    short = await client.post(f"{API}/decisions", json={**payload, "idempotency_key": "abc"})
    assert short.status_code == 422


async def test_decision_records_the_alternative(client, created_decision_ids):
    """alternative_org_code is derived from recommendation_id and returned with the alternative's name."""
    alerts = (await client.get(f"{API}/alerts", params={"limit": 60})).json()["items"]
    for alert in alerts:
        recs = (
            await client.get(f"{API}/hospitals/{alert['org_code']}/profiles/{alert['profile_code']}/recommendations")
        ).json()
        if recs["alternatives"]:
            break
    else:
        pytest.fail("no alert with recommendation alternatives")
    alternative = recs["alternatives"][0]
    payload = decision_payload(alert, action="confirm", recommendation_id=alternative["recommendation_id"])
    created = await client.post(f"{API}/decisions", json=payload)
    assert created.status_code == 201, created.text
    body = created.json()
    created_decision_ids.append(body["id"])
    assert body["alternative_org_code"] == alternative["org_code"]
    assert body["alternative_org_name"] == alternative["org_name"]

    listed = (await client.get(f"{API}/decisions", params={"org": alert["org_code"], "limit": 500})).json()["items"]
    assert next(d for d in listed if d["id"] == body["id"])["alternative_org_name"] == alternative["org_name"]

    mismatch = {**payload, "alternative_org_code": alert["org_code"], "actor": f"{TEST_ACTOR_PREFIX}mismatch"}
    assert (await client.post(f"{API}/decisions", json=mismatch)).status_code == 422


# ------------------------------------------------------------------------------------------ export
async def test_export_xlsx(client, high_load):
    org, profile = high_load["org_code"], high_load["profile_code"]
    response = await client.get(f"{API}/hospitals/{org}/profiles/{profile}/export", params={"format": "xlsx"})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response.headers["content-disposition"].startswith("attachment;")
    assert f"{org}_{profile}" in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    assert workbook.sheetnames == ["Карточка", "Ряд по дням", "Прогноз 14 дней", "Почему", "Рекомендации", "Решения"]
    card = (await client.get(f"{API}/hospitals/{org}/profiles/{profile}")).json()
    series_rows = list(workbook["Ряд по дням"].iter_rows(min_row=2, values_only=True))
    assert len(series_rows) == len(card["series"])
    assert series_rows[-1][4] == card["status"]["queue_now"]
    first_sheet = [c for row in workbook["Карточка"].iter_rows(values_only=True) for c in row if c is not None]
    assert card["status"]["org_name"] in " ".join(map(str, first_sheet))


async def test_export_pdf(client, high_load):
    org, profile = high_load["org_code"], high_load["profile_code"]
    response = await client.get(f"{API}/hospitals/{org}/profiles/{profile}/export", params={"format": "pdf"})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF-") and len(response.content) > 5_000
    assert (
        await client.get(f"{API}/hospitals/{org}/profiles/{profile}/export", params={"format": "docx"})
    ).status_code == 422
    assert (await client.get(f"{API}/hospitals/NOPE/profiles/{profile}/export")).status_code == 404


async def test_viewer_may_export(anon_client, api_keys, high_load):
    path = f"{API}/hospitals/{high_load['org_code']}/profiles/{high_load['profile_code']}/export"
    assert (await anon_client.get(path, headers=auth(api_keys["viewer"]))).status_code == 200
    assert (await anon_client.get(path)).status_code == 401


# ------------------------------------------------------------------------------------------ config, filters
async def test_config(client):
    body = (await client.get(f"{API}/config")).json()
    weights = body["load_index"]["weights"]
    assert sum(weights.values()) == pytest.approx(1.0)
    assert body["load_index"]["refusal_rate_cap"] > 0 and body["load_index"]["queue_trend_cap_pct"] > 0
    assert body["status_thresholds"]["high"] > body["status_thresholds"]["elevated"]
    assert body["alerts"]["load_index_min"] == body["status_thresholds"]["high"]
    source = body["data_source"]
    assert source["publisher"] and source["description"] and source["datasets"]
    overview = (await client.get(f"{API}/overview")).json()
    assert body["as_of_date"] == overview["as_of_date"]


async def test_alert_filters(client, high_load):
    high = (await client.get(f"{API}/alerts", params={"status": "high", "limit": 500})).json()
    assert high["items"] and all(
        a["status"] == "high" and a["status_label"] == "Высокая нагрузка" for a in high["items"]
    )
    everything = (await client.get(f"{API}/alerts", params={"limit": 1})).json()["total"]
    assert high["total"] <= everything

    profile = high_load["profile_code"]
    by_profile = (await client.get(f"{API}/alerts", params={"profile": profile, "limit": 500})).json()
    assert by_profile["items"] and all(a["profile_code"] == profile for a in by_profile["items"])
    assert (await client.get(f"{API}/alerts", params={"status": "busy"})).status_code == 422
    assert (await client.get(f"{API}/alerts", params={"profile": "nope"})).status_code == 404


async def test_referrals_have_diagnosis_names(client, high_load):
    body = (
        await client.get(
            f"{API}/hospitals/{high_load['org_code']}/profiles/{high_load['profile_code']}/referrals",
            params={"limit": 50},
        )
    ).json()
    coded = [r for r in body["items"] if r["icd10_code"]]
    assert coded and all(r["diagnosis_name"] for r in coded), "every coded diagnosis has a name in dim_icd"


async def test_cors_preflight_needs_no_key(anon_client: httpx.AsyncClient):
    response = await anon_client.options(
        f"{API}/overview",
        headers={
            "Origin": "http://localhost:5173",
            "Access-Control-Request-Method": "GET",
            "Access-Control-Request-Headers": "x-api-key",
        },
    )
    assert response.status_code == 200
    assert "x-api-key" in response.headers.get("access-control-allow-headers", "").lower()
