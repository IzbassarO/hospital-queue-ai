"""Hospital × profile card as a file: XLSX (openpyxl) or PDF (reportlab).

Same content as the UI card: status and load_index components, KPIs, daily series, 14-day forecast (no queue
forecast), explanation factors, recommendations with the disclaimer, and the decision history. Built from the
same service functions as the JSON endpoints, so the numbers match the screen.
"""

import datetime as dt
import io
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.lineplots import LinePlot
from reportlab.graphics.shapes import Drawing, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.schemas.activity import Decision, RecommendationResponse
from app.schemas.status import HospitalProfileCard
from app.services import activity, recommend, status

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MEDIA_TYPE = "application/pdf"
STATUS_SHORT = {
    "high": "Высокая",
    "elevated": "Повышенная",
    "normal": "Норма",
    "insufficient_data": "Недостаточно данных",
}
ACTIONS = {"confirm": "Подтверждено", "reject": "Отклонено", "defer": "Отложено"}
DISCLAIMER_HUMAN = "Решение принимает специалист. Система только предлагает."


class ExportUnavailableError(Exception):
    """The export cannot be produced on this server (e.g. no Cyrillic font for PDF); mapped to HTTP 503."""


@dataclass(frozen=True)
class ExportFile:
    content: bytes
    media_type: str
    filename: str


@dataclass(frozen=True)
class CardData:
    card: HospitalProfileCard
    recommendations: RecommendationResponse
    decisions: list[Decision]


def _load(session: Session, org_code: str, profile_code: str) -> CardData:
    card = status.hospital_card(session, org_code, profile_code)
    recs = recommend.recommend(session, org_code, profile_code)
    decisions = activity.list_decisions(session, org_code, profile_code, limit=500, offset=0).items
    return CardData(card=card, recommendations=recs, decisions=decisions)


# ------------------------------------------------------------------------------------------ formatting
def _num(value: float | int | None, digits: int = 0, suffix: str = "") -> str:
    if value is None:
        return "—"
    text = f"{value:,.{digits}f}".replace(",", " ").replace(".", ",")
    return f"{text}{suffix}"


def _pct(share: float | None) -> str:
    return "—" if share is None else _num(share * 100, 1, "%")


def _date(value: dt.date | dt.datetime | None) -> str:
    return "—" if value is None else value.strftime("%d.%m.%Y %H:%M" if isinstance(value, dt.datetime) else "%d.%m.%Y")


def _kpis(card: HospitalProfileCard) -> list[tuple[str, str]]:
    s = card.status
    return [
        ("Индекс нагрузки", _num(s.load_index, 1)),
        ("Статус", STATUS_SHORT[s.status]),
        (
            "Место в регионе",
            f"{s.region_rank} из {s.region_n_ranked}" if s.region_rank is not None else "—",
        ),
        ("Очередь сейчас", _num(s.queue_now)),
        ("Срок рассасывания очереди", _num(s.backlog_days, 1, " дн.")),
        ("Медиана ожидания (28 дней)", _num(s.median_wait_28d, 1, " дн.")),
        ("Доля отказов (28 дней)", _pct(s.refusal_rate_28d)),
        ("Направления за 28 дней", _num(s.registrations_28d)),
        ("Прогноз направлений на 14 дней", _num(s.forecast_registrations_14d)),
        ("Прогноз госпитализаций на 14 дней", _num(s.forecast_hospitalizations_14d)),
        ("Рост очереди сверх медианы по стране", _num(s.queue_trend_4w, 1, " п.п./нед.")),
        ("Составляющая: срок рассасывания", _num(s.components.backlog_score, 2)),
        ("Составляющая: доля отказов", _num(s.components.refusal_score, 2)),
        ("Составляющая: рост очереди", _num(s.components.trend_score, 2)),
    ]


def _filename(card: HospitalProfileCard, ext: str) -> str:
    s = card.status
    return f"hqai_card_{s.org_code}_{s.profile_code}_{s.as_of_date.isoformat()}.{ext}"


# ------------------------------------------------------------------------------------------ XLSX
_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="EEF1F5")


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list], widths: list[int]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, dt.datetime):
                cell.number_format = "DD.MM.YYYY HH:MM"
            elif isinstance(cell.value, dt.date):
                cell.number_format = "DD.MM.YYYY"
            elif isinstance(cell.value, str) and len(cell.value) > 60:
                cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_xlsx(data: CardData) -> bytes:
    card, recs = data.card, data.recommendations
    s = card.status
    wb = Workbook()
    ws = wb.active
    ws.title = "Карточка"
    ws.append(["Карточка стационара по профилю"])
    ws["A1"].font = Font(bold=True, size=14)
    for label, value in (
        ("Стационар", f"{s.org_name} ({s.org_code})"),
        ("Регион", f"{s.region_name} ({s.region_code})"),
        ("Профиль", f"{s.profile_name} ({s.profile_code})"),
        ("Данные на", s.as_of_date),
        ("Сформировано", dt.datetime.now().replace(microsecond=0)),
    ):
        ws.append([label, value])
    ws.append([])
    ws.append(["Показатель", "Значение"])
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
    for label, value in _kpis(card):
        ws.append([label, value])
    ws.append([])
    ws.append([DISCLAIMER_HUMAN])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 90
    for row in ws.iter_rows(min_row=2, max_row=6):
        if isinstance(row[1].value, dt.datetime):
            row[1].number_format = "DD.MM.YYYY HH:MM"
        elif isinstance(row[1].value, dt.date):
            row[1].number_format = "DD.MM.YYYY"

    _sheet(
        wb,
        "Ряд по дням",
        ["Дата", "Направления", "Госпитализации", "Отказы", "Очередь"],
        [[p.date, p.registrations, p.hospitalizations, p.refusals, p.queue] for p in card.series],
        [14, 14, 16, 10, 10],
    )
    forecast_rows = [[p.date, p.horizon, p.registrations, p.hospitalizations] for p in card.forecast.points]
    forecast_rows += [[], [card.forecast.note]]
    _sheet(
        wb,
        "Прогноз 14 дней",
        ["Дата", "Горизонт, дней", "Направления (прогноз)", "Госпитализации (прогноз)"],
        forecast_rows,
        [14, 16, 24, 26],
    )
    factor_rows = []
    for title, factors in (
        ("Время ожидания", card.explanation_factors.wait_time),
        ("Риск отказа", card.explanation_factors.refusal_risk),
    ):
        for f in factors:
            factor_rows.append(
                [title, f.short_label, f.most_common_value_display, f.mean_effect, f.unit, f.share_in_top5]
            )
    _sheet(
        wb,
        "Почему",
        ["Модель", "Фактор", "Типичное значение", "Среднее влияние", "Единица", "Доля направлений с фактором в топ-5"],
        factor_rows,
        [18, 34, 60, 16, 10, 22],
    )
    rec_rows = [
        [
            a.org_name,
            a.org_code,
            a.expected_wait_current,
            a.expected_wait_alternative,
            a.delta_days,
            a.refusal_rate_current,
            a.refusal_rate_alternative,
            a.backlog_days_current,
            a.backlog_days_alternative,
            a.explanation,
        ]
        for a in recs.alternatives
    ]
    if not recs.alternatives and recs.reason:
        rec_rows.append([recs.reason])
    rec_rows += [[], ["Метод: оценка по историческим медианам"], [recs.disclaimer], [DISCLAIMER_HUMAN]]
    _sheet(
        wb,
        "Рекомендации",
        [
            "Альтернатива",
            "Код",
            "Ожидание сейчас, дн.",
            "Ожидание в альтернативе, дн.",
            "Разница, дн.",
            "Отказы сейчас",
            "Отказы в альтернативе",
            "Рассасывание сейчас, дн.",
            "Рассасывание в альтернативе, дн.",
            "Пояснение",
        ],
        rec_rows,
        [50, 8, 12, 14, 10, 12, 12, 14, 16, 90],
    )
    _sheet(
        wb,
        "Решения",
        ["Дата и время", "Решение", "Альтернатива", "Кто", "Комментарий", "API-ключ"],
        [
            [
                d.created_at.replace(tzinfo=None),
                ACTIONS[d.action],
                d.alternative_org_name or d.alternative_org_code or "",
                d.actor,
                d.comment or "",
                d.api_key_label or "",
            ]
            for d in data.decisions
        ],
        [18, 14, 50, 30, 50, 24],
    )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------------------------------ PDF
@lru_cache
def _fonts() -> tuple[str, str]:
    settings = get_settings()
    regular = next((p for p in settings.pdf_font_paths if Path(p).is_file()), None)
    if regular is None:
        raise ExportUnavailableError(
            "PDF export needs a TrueType font with Cyrillic glyphs; none of PDF_FONT_PATHS exists on this server"
        )
    bold = next((p for p in settings.pdf_bold_font_paths if Path(p).is_file()), regular)
    pdfmetrics.registerFont(TTFont("HqaiSans", regular))
    pdfmetrics.registerFont(TTFont("HqaiSans-Bold", bold))
    return "HqaiSans", "HqaiSans-Bold"


def _line_chart(title: str, series: list[tuple[str, list[tuple[float, float]], colors.Color, bool]], font: str):
    width, height = 170 * mm, 55 * mm
    drawing = Drawing(width, height)
    plot = LinePlot()
    plot.x, plot.y, plot.width, plot.height = 30, 20, width - 45, height - 40
    plot.data = [points for _, points, _, _ in series]
    for i, (_, _, color, dashed) in enumerate(series):
        plot.lines[i].strokeColor = color
        plot.lines[i].strokeWidth = 1.4
        if dashed:
            plot.lines[i].strokeDashArray = [4, 3]
    plot.xValueAxis.labels.fontName = plot.yValueAxis.labels.fontName = font
    plot.xValueAxis.labels.fontSize = plot.yValueAxis.labels.fontSize = 7
    plot.xValueAxis.labelTextFormat = lambda ordinal: dt.date.fromordinal(int(ordinal)).strftime("%d.%m")
    plot.yValueAxis.valueMin = 0
    drawing.add(plot)
    drawing.add(String(30, height - 10, title, fontName=font, fontSize=9))
    x = width - 10
    for label, _, color, dashed in reversed(series):
        text = f"{label}{' (пунктир)' if dashed else ''}"
        x -= pdfmetrics.stringWidth(text, font, 7) + 14
        drawing.add(String(x, height - 10, text, fontName=font, fontSize=7, fillColor=color))
    return drawing


def build_pdf(data: CardData) -> bytes:
    font, bold = _fonts()
    card, recs = data.card, data.recommendations
    s = card.status
    base = ParagraphStyle("base", fontName=font, fontSize=9, leading=12)
    small = ParagraphStyle("small", parent=base, fontSize=8, leading=10, textColor=colors.HexColor("#536072"))
    h1 = ParagraphStyle("h1", parent=base, fontName=bold, fontSize=14, leading=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=base, fontName=bold, fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)
    grid = TableStyle(
        [
            ("FONT", (0, 0), (-1, -1), font, 8),
            ("FONT", (0, 0), (-1, 0), bold, 8),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF1F5")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D5DBE3")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]
    )

    def table(rows: list[list], widths: list[float], right_from: int | None = None) -> Table:
        wrapped = [[Paragraph(str(c), base) if isinstance(c, str) and len(c) > 40 else c for c in r] for r in rows]
        t = Table(wrapped, colWidths=[w * mm for w in widths], repeatRows=1)
        t.setStyle(grid)
        if right_from is not None:
            t.setStyle(TableStyle([("ALIGN", (right_from, 1), (-1, -1), "RIGHT")]))
        return t

    story: list = [
        Paragraph(s.org_name, h1),
        Paragraph(
            f"{s.region_name} · профиль «{s.profile_name}» ({s.profile_code}) · данные на {_date(s.as_of_date)}", base
        ),
        Paragraph(f"Сформировано {_date(dt.datetime.now())}. {DISCLAIMER_HUMAN}", small),
        Paragraph("Ключевые показатели", h2),
        table([["Показатель", "Значение"], *[[k, v] for k, v in _kpis(card)]], [95, 75], right_from=1),
    ]

    observed = [(p.date.toordinal(), p.queue) for p in card.series]
    story += [
        Paragraph("Динамика", h2),
        _line_chart("Очередь по дням", [("Очередь", observed, colors.HexColor("#1c2430"), False)], font),
    ]
    regs = [(p.date.toordinal(), p.registrations) for p in card.series]
    fc = [(p.date.toordinal(), p.registrations) for p in card.forecast.points]
    chart_series = [("Направления, факт", regs, colors.HexColor("#1f5fae"), False)]
    if fc:
        chart_series.append(("прогноз", fc, colors.HexColor("#1f5fae"), True))
    story += [
        _line_chart("Направления: факт и прогноз на 14 дней", chart_series, font),
        Paragraph(card.forecast.note, small),
        Paragraph("Прогноз на 14 дней", h2),
        table(
            [
                ["Дата", "Направления", "Госпитализации"],
                *[[_date(p.date), _num(p.registrations, 1), _num(p.hospitalizations, 1)] for p in card.forecast.points],
            ],
            [40, 40, 40],
            right_from=1,
        ),
        Paragraph("Почему (факторы моделей, связи в данных, а не причины)", h2),
    ]
    factor_rows = [["Модель", "Фактор", "Типичное значение", "Влияние"]]
    for title, factors in (
        ("Время ожидания", card.explanation_factors.wait_time),
        ("Риск отказа", card.explanation_factors.refusal_risk),
    ):
        factor_rows += [
            [title, f.short_label, f.most_common_value_display, _num(f.mean_effect, 1, f" {f.unit}")] for f in factors
        ]
    story.append(table(factor_rows, [28, 42, 80, 20], right_from=3))

    story.append(Paragraph("Рекомендации (оценка по историческим медианам)", h2))
    if recs.alternatives:
        story.append(
            table(
                [
                    ["Альтернатива", "Ожидание: сейчас → альт.", "Разница", "Отказы: сейчас → альт."],
                    *[
                        [
                            a.org_name,
                            f"{_num(a.expected_wait_current, 1)} → {_num(a.expected_wait_alternative, 1)} дн.",
                            _num(a.delta_days, 1, " дн."),
                            f"{_pct(a.refusal_rate_current)} → {_pct(a.refusal_rate_alternative)}",
                        ]
                        for a in recs.alternatives
                    ],
                ],
                [80, 38, 18, 34],
            )
        )
        for a in recs.alternatives:
            story.append(Paragraph(a.explanation, small))
    else:
        story.append(Paragraph(recs.reason or "Рекомендации не формируются.", base))
    story += [Paragraph(recs.disclaimer, small), Paragraph("История решений", h2)]
    if data.decisions:
        story.append(
            table(
                [
                    ["Дата и время", "Решение", "Альтернатива", "Кто", "Комментарий"],
                    *[
                        [
                            _date(d.created_at),
                            ACTIONS[d.action],
                            d.alternative_org_name or d.alternative_org_code or "—",
                            d.actor,
                            d.comment or "—",
                        ]
                        for d in data.decisions
                    ],
                ],
                [26, 22, 50, 32, 40],
            )
        )
    else:
        story.append(Paragraph("Решений пока нет.", base))

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont(font, 7)
        canvas.setFillColor(colors.HexColor("#536072"))
        canvas.drawString(15 * mm, 10 * mm, f"hospital-queue-ai · {s.org_code} × {s.profile_code} · {DISCLAIMER_HUMAN}")
        canvas.drawRightString(A4[0] - 15 * mm, 10 * mm, f"стр. {doc.page}")
        canvas.restoreState()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=18 * mm,
        title=f"Карточка {s.org_code} × {s.profile_code}",
        author="hospital-queue-ai",
    )
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def export_card(session: Session, org_code: str, profile_code: str, fmt: str) -> ExportFile:
    data = _load(session, org_code, profile_code)
    if fmt == "xlsx":
        return ExportFile(build_xlsx(data), XLSX_MEDIA_TYPE, _filename(data.card, "xlsx"))
    return ExportFile(build_pdf(data), PDF_MEDIA_TYPE, _filename(data.card, "pdf"))
